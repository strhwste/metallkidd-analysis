"""Liest metallkidd.xlsx (Zenodo 11506192) einmal ein und legt Formeln und
Cache-Werte als wb.pkl ab. verifikation.py braucht diese Datei."""
import pickle
import openpyxl

formeln = openpyxl.load_workbook("metallkidd.xlsx")
werte = openpyxl.load_workbook("metallkidd.xlsx", data_only=True)
lies = lambda wb: {ws.title: {c.coordinate: c.value for r in ws.iter_rows() for c in r if c.value is not None}
                   for ws in wb.worksheets}
pickle.dump((lies(formeln), lies(werte)), open("wb.pkl", "wb"))
print("wb.pkl geschrieben")
